import sys
from pathlib import Path
import openpyxl
import pytest

# Add backend/scripts to sys.path so we can import TEMPLATES and generator functions
backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))
sys.path.insert(0, str(backend_dir / "scripts"))

from scripts.generate_standard_templates import (
    TEMPLATES,
    ColumnType,
    build_workbook_for_spec,
    generate_empty_templates,
    verify_generated_templates,
)


@pytest.fixture(scope="module")
def templates_dir(tmp_path_factory):
    temp_dir = tmp_path_factory.mktemp("templates")
    generate_empty_templates(temp_dir)
    return temp_dir


def test_all_five_templates_generated(templates_dir):
    """Verifies that all 5 business-approved templates are generated."""
    expected_files = [spec.file_name for spec in TEMPLATES.values()]
    assert len(expected_files) == 5

    for file_name in expected_files:
        target = templates_dir / file_name
        assert target.exists(), f"Missing template file: {file_name}"
        assert target.stat().st_size > 0, f"Template file is empty: {file_name}"


def test_tenders_template_spec(templates_dir):
    """Validates 01_Реестр_Тендеров_и_Контрактов.xlsx headers and validations."""
    file_path = templates_dir / "01_Реестр_Тендеров_и_Контрактов.xlsx"
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    assert ws.title == "Реестр_Тендеров"
    assert ws.freeze_panes == "A2"

    headers = [cell.value for cell in ws[1]]
    expected_headers = [
        "Номер_Закупки",
        "Наименование_Объекта",
        "Заказчик",
        "Сумма_Контракта_Руб",
        "Сумма_Обеспечения_Заявки",
        "Статус_Тендера",
        "Дедлайн_Подачи",
        "Дата_Окончания_Работ",
        "Сумма_Закрыто_КС_Руб",
        "Оплачено_Заказчиком_Руб",
        "Куратор_Ответственный",
        "Банковская_Гарантия_До",
    ]
    assert headers == expected_headers

    # Validate DataValidation
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 1
    dv = dvs[0]
    assert dv.type == "list"
    expected_options = "Подготовка,Подано,Торги,Победа,Контракт,Исполнено,Отклонено"
    assert expected_options in dv.formula1
    wb.close()


def test_cashflow_template_spec(templates_dir):
    """Validates 02_ДДС_Касса_и_Взаиморасчеты.xlsx headers and validations."""
    file_path = templates_dir / "02_ДДС_Касса_и_Взаиморасчеты.xlsx"
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    assert ws.title == "ДДС_Касса"
    assert ws.freeze_panes == "A2"

    headers = [cell.value for cell in ws[1]]
    expected_headers = [
        "Дата_Операции",
        "Тип_Операции",
        "Статья_ДДС",
        "Контрагент",
        "Сумма_Руб",
        "Статус_Платежа",
        "Расчетный_Счет",
        "Основание_Платежа",
    ]
    assert headers == expected_headers

    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 2
    formulas = [dv.formula1 for dv in dvs]
    assert any("Поступление,Списание" in f for f in formulas)
    assert any("Факт_Исполнено,План_Ожидается" in f for f in formulas)
    wb.close()


def test_procurement_template_spec(templates_dir):
    """Validates 03_План_Закупок_и_Материалов.xlsx headers and validations."""
    file_path = templates_dir / "03_План_Закупок_и_Материалов.xlsx"
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    assert ws.title == "План_Закупок"
    assert ws.freeze_panes == "A2"

    headers = [cell.value for cell in ws[1]]
    expected_headers = [
        "Объект_Привязка",
        "Наименование_ТМЦ",
        "Количество_Ед",
        "Единица_Изм",
        "Поставщик",
        "Сумма_Счета_Руб",
        "Статус_Снабжения",
        "Срок_Доставки_Дедлайн",
    ]
    assert headers == expected_headers

    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 1
    dv = dvs[0]
    expected_options = "Заявка,Счет_На_Оплате,Оплачено,В_Пути,На_Объекте"
    assert expected_options in dv.formula1
    wb.close()


def test_dispatch_template_spec(templates_dir):
    """Validates 04_Диспетчерский_График_и_Техника.xlsx headers and validations."""
    file_path = templates_dir / "04_Диспетчерский_График_и_Техника.xlsx"
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    assert ws.title == "Диспетчерский_График"
    assert ws.freeze_panes == "A2"

    headers = [cell.value for cell in ws[1]]
    expected_headers = [
        "Дата_Смены",
        "Объект_Назначения",
        "Тип_Ресурса",
        "Наименование_Единицы",
        "Ответственный_Водитель_Бригадир",
        "Статус_Работы",
        "Замечания_Инциденты",
    ]
    assert headers == expected_headers

    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 2
    formulas = [dv.formula1 for dv in dvs]
    assert any("Спецтехника,Бригада_Рабочих,Инженер_Куратор" in f for f in formulas)
    assert any("На_Объекте,В_Ремонте,Простой_Авария,Завершено" in f for f in formulas)
    wb.close()


def test_tasks_template_spec(templates_dir):
    """Validates 05_Операционные_Задачи_и_Алерты.xlsx headers and validations."""
    file_path = templates_dir / "05_Операционные_Задачи_и_Алерты.xlsx"
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    assert ws.title == "Задачи_и_Алерты"
    assert ws.freeze_panes == "A2"

    headers = [cell.value for cell in ws[1]]
    expected_headers = [
        "Заголовок_Задачи",
        "Связанный_Контракт",
        "Ответственная_Роль",
        "ФИО_Исполнителя",
        "Крайний_Срок_Дедлайн",
        "Приоритет_Срочность",
        "Статус_Исполнения",
        "Потенциальный_Штраф_Руб",
    ]
    assert headers == expected_headers

    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 3
    formulas = [dv.formula1 for dv in dvs]
    assert any("Гендиректор,Замгендиректора,Тендерщик" in f for f in formulas)
    assert any("Обычный,Высокий,КРИТИЧЕСКИЙ_ШТРАФ" in f for f in formulas)
    assert any("В_Работе,На_Проверке,Закрыто" in f for f in formulas)
    wb.close()


def test_verifier_function(templates_dir):
    """Ensures verify_generated_templates returns True for valid directory."""
    assert verify_generated_templates(templates_dir) is True
