import os
import tempfile
from datetime import date
from decimal import Decimal
import openpyxl
import pytest

from app.models.enums import DomainCategory
from app.services.parser import ExcelParserService


@pytest.fixture
def sample_merged_excel_file():
    """
    Creates a temporary Excel file with multi-row merged cells,
    merged data cells, decorative banners, and total summary row.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Реестр договоров"

    # Row 1: Decorative banner / multi-column merged title
    ws["A1"] = "Отчет по строительно-монтажным работам филиала"
    ws.merge_cells("A1:E1")

    # Row 2: Secondary merged header categories
    ws["A2"] = "Основные параметры"
    ws.merge_cells("A2:B2")
    ws["C2"] = "Финансы и сроки"
    ws.merge_cells("C2:E2")

    # Row 3: Actual column headers
    ws["A3"] = "Объект строительства"
    ws["B3"] = "Куратор от заказчика"
    ws["C3"] = "Цена контракта, руб."
    ws["D3"] = "Закрыто актами КС"
    ws["E3"] = "Срок сдачи объекта"

    # Row 4: Data row 1
    ws["A4"] = "Кислородный цех №101"
    ws["B4"] = "Петров В.А."
    ws["C4"] = "12 500 000,00"
    ws["D4"] = "8 200 000,00"
    ws["E4"] = "15.10.2026"

    # Row 5 & 6: Data row with merged cell across 2 rows for project name
    ws["A5"] = "Котельная БМК-12"
    ws["B5"] = "Сидоров М.И."
    ws["C5"] = "4 800 000,00"
    ws["D5"] = "4 800 000,00"
    ws["E5"] = "01.11.2026"

    ws["A6"] = None  # merged with A5
    ws["B6"] = "Иванов И.И. (сменщик)"
    ws["C6"] = "1 200 000,00"
    ws["D6"] = "600 000,00"
    ws["E6"] = "10.11.2026"
    ws.merge_cells("A5:A6")

    # Row 7: Total Summary row (must be excluded from parsed records)
    ws["A7"] = "Итого по всем объектам:"
    ws["B7"] = ""
    ws["C7"] = "18 500 000,00"
    ws["D7"] = "13 600 000,00"
    ws["E7"] = ""

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name

    wb.save(tmp_path)
    wb.close()

    yield tmp_path

    if os.path.exists(tmp_path):
        os.remove(tmp_path)


def test_excel_unmerge_and_fill(sample_merged_excel_file):
    """
    Validates Step 1 of the parser pipeline:
    Merged cell A5:A6 must be unmerged and A6 populated with value from A5.
    """
    wb = openpyxl.load_workbook(sample_merged_excel_file)
    ws = wb.active

    # Pre-condition: A6 is empty and merged with A5
    assert len(ws.merged_cells.ranges) > 0

    ExcelParserService.unmerge_and_normalize_sheet(ws)

    # Post-condition: no merged cells remain and value is propagated
    assert len(ws.merged_cells.ranges) == 0
    assert ws["A5"].value == "Котельная БМК-12"
    assert ws["A6"].value == "Котельная БМК-12"
    wb.close()


def test_detect_table_structure_and_exclude_summary(sample_merged_excel_file):
    """
    Validates header detection and automatic exclusion of summary/total rows.
    """
    wb = openpyxl.load_workbook(sample_merged_excel_file)
    ws = wb.active

    header_row_idx, headers, data_rows = ExcelParserService.detect_table_structure(ws)
    wb.close()

    assert header_row_idx == 3
    assert "Объект строительства" in headers
    assert "Цена контракта, руб." in headers
    assert "Срок сдачи объекта" in headers

    # Exactly 3 operational data rows (row 4, 5, 6). Row 7 'Итого...' must be filtered out!
    assert len(data_rows) == 3
    # Check that A6 inherited project name from A5
    assert data_rows[2][0] == "Котельная БМК-12"
    # Ensure no summary row in data
    assert not any("итого" in str(r[0]).lower() for r in data_rows)


def test_calculate_headers_fingerprint():
    """Validates deterministic SHA-256 fingerprint generation."""
    headers_1 = ["Объект строительства", "Цена контракта, руб.", "Срок сдачи"]
    headers_2 = [" Объект строительства ", "Цена контракта, руб.", "Срок сдачи "]

    fp_1 = ExcelParserService.calculate_headers_fingerprint(headers_1)
    fp_2 = ExcelParserService.calculate_headers_fingerprint(headers_2)

    assert fp_1 == fp_2
    assert len(fp_1) == 64


def test_infer_domain_and_mapping(sample_merged_excel_file):
    """
    Validates heuristic semantic mapping to canonical Domain 1 (Tenders & Production).
    """
    wb = openpyxl.load_workbook(sample_merged_excel_file)
    ws = wb.active
    _, headers, data_rows = ExcelParserService.detect_table_structure(ws)
    wb.close()

    domain, confidence, mapping = ExcelParserService.infer_domain_and_mapping(
        headers, data_rows[:3]
    )

    assert domain == DomainCategory.PRODUCTION_TENDERS
    assert confidence >= 0.80

    mapping_dict = {m["source_column"]: m["target_field"] for m in mapping}
    assert mapping_dict.get("Объект строительства") == "project_name"
    assert mapping_dict.get("Цена контракта, руб.") == "contract_sum"
    assert mapping_dict.get("Срок сдачи объекта") == "deadline_date"
    assert mapping_dict.get("Куратор от заказчика") == "curator_name"


def test_clean_decimal_and_date_helpers():
    """Validates cleaning of messy Russian currency strings and multiple date formats."""
    assert ExcelParserService.clean_decimal(" 12 500 000,50 руб. ") == Decimal("12500000.50")
    assert ExcelParserService.clean_decimal("4.800.000,00 ₽") == Decimal("4800000.00")
    assert ExcelParserService.clean_decimal(None) == Decimal("0.00")
    assert ExcelParserService.clean_decimal("invalid") == Decimal("0.00")

    assert ExcelParserService.clean_date("15.10.2026") == date(2026, 10, 15)
    assert ExcelParserService.clean_date("2026-11-01") == date(2026, 11, 1)
    assert ExcelParserService.clean_date(None) is None
