import io
import openpyxl
import pytest
from app.services.customer_tenders import read_customer_workbook

def workbook():
    w=openpyxl.Workbook()
    a=w.active
    a.title='Реест Ольга'
    a.append(['вн.№','Заказчик']+[None]*20)
    a.append([None]*12+['НМЦК','Процент падения','Сумма падения','Сумма контракта (выиграли)'])
    a.append([None]*22)
    a.append([2501,'Заказчик']+[None]*13+[0])
    b=w.create_sheet('Ежедневный мониторинг')
    b.append(['вн.№','Заказчик','Юрлицо','Город']+['Дата','Документ','Ссылка','ЕИС','Подписано','Оплата','Полностью']*3)
    b.append([None]*25)
    b.append([None]*4+['окт.26']+[None]*6+['ноя.26']+[None]*6+['дек.26'])
    b.append(['2501','Заказчик',None,None,'2026-10-26']+[None]*20)
    b.append(['#REF!','#REF!'])
    f=io.BytesIO();w.save(f);return f.getvalue()

def test_two_sheets_and_formula_errors():
    d=read_customer_workbook(workbook())
    assert len(d['Реест Ольга']['rows'])==1
    assert len(d['Ежедневный мониторинг']['rows'])==1
    assert d['Реест Ольга']['rows'][0]['key']==d['Ежедневный мониторинг']['rows'][0]['key']=='2501'
    assert d['Реест Ольга']['rows'][0]['values'][15]==0
    assert len(d['Ежедневный мониторинг']['warnings'])==1
    assert 'Период 2' in d['Ежедневный мониторинг']['columns'][11]
    assert len(d['Ежедневный мониторинг']['rows'][0]['values'])==25

def test_other_template_is_not_customer_register():
    w=openpyxl.Workbook();f=io.BytesIO();w.save(f)
    assert read_customer_workbook(f.getvalue()) is None
