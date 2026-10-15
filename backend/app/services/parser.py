import re
import uuid
import hashlib
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from app.models.enums import DomainCategory

# -----------------------------------------------------------------------------
# Canonical Domains Ontology & Synonyms Dictionary
# -----------------------------------------------------------------------------

CANONICAL_FIELDS: Dict[DomainCategory, List[Dict[str, Any]]] = {
    DomainCategory.PRODUCTION_TENDERS: [
        {"field": "project_name", "label": "Название объекта / проекта", "required": True, "synonyms": ["объект", "наименование", "проект", "стройка", "титул", "название", "объект смр", "объект строительства", "контракт"]},
        {"field": "project_code", "label": "Номер / шифр контракта", "required": False, "synonyms": ["номер", "шифр", "код", "№", "№ контракта", "№ договора", "артикул", "код проекта"]},
        {"field": "contract_sum", "label": "Сумма договора (руб.)", "required": True, "synonyms": ["цена контракта", "договорная цена", "сумма договора", "стоимость", "цена", "сумма", "цена с ндс", "бюджет", "контрактная цена"]},
        {"field": "completed_sum", "label": "Выполнено / Закрыто по КС (руб.)", "required": False, "synonyms": ["закрыто", "выполнено", "кс-2", "кс-3", "акты", "объем выполнения", "закрыто по кс", "закрыто актами кс"]},
        {"field": "paid_sum", "label": "Оплачено заказчиком (руб.)", "required": False, "synonyms": ["оплачено", "платеж", "поступило", "факт оплаты", "оплата"]},
        {"field": "debt_sum", "label": "Остаток долга (руб.)", "required": False, "synonyms": ["остаток", "долг", "задолженность", "дебиторка", "остаток к оплате", "к оплате"]},
        {"field": "deadline_date", "label": "Срок сдачи / Дедлайн", "required": True, "synonyms": ["срок сдачи", "срок выполнения", "дедлайн", "дата окончания", "срок", "дата сдачи объекта", "срок работ"]},
        {"field": "curator_name", "label": "Ответственный куратор", "required": False, "synonyms": ["куратор", "прораб", "ответственный", "менеджер", "куратор от заказчика", "руководитель"]},
        {"field": "customer_name", "label": "Заказчик / Генподрядчик", "required": False, "synonyms": ["заказчик", "клиент", "генподрядчик", "контрагент"]},
    ],
    DomainCategory.CASHFLOW_DEBT: [
        {"field": "operation_date", "label": "Дата операции", "required": True, "synonyms": ["дата", "число", "дата платежа", "дата операции", "период"]},
        {"field": "counterparty", "label": "Контрагент", "required": False, "synonyms": ["контрагент", "получатель", "плательщик", "поставщик", "клиент", "партнер"]},
        {"field": "direction", "label": "Направление (приход/расход)", "required": True, "synonyms": ["направление", "тип", "вид", "приход", "расход", "вид движения"]},
        {"field": "actual_amount", "label": "Фактическая сумма", "required": True, "synonyms": ["сумма", "факт", "фактическая сумма", "сумма платежа", "объем"]},
        {"field": "planned_amount", "label": "Плановая сумма", "required": False, "synonyms": ["план", "плановая сумма", "ожидаемо"]},
        {"field": "category", "label": "Статья расходов / доходов", "required": False, "synonyms": ["статья", "категория", "назначение", "вид расхода", "ддс"]},
        {"field": "payment_due_date", "label": "Срок оплаты", "required": False, "synonyms": ["срок оплаты", "дедлайн оплаты", "плановый срок"]},
        {"field": "account_name", "label": "Банковский счет / Касса", "required": False, "synonyms": ["счет", "банк", "р/с", "касса"]},
    ],
    DomainCategory.PAYROLL_HR: [
        {"field": "employee_name", "label": "ФИО сотрудника", "required": True, "synonyms": ["сотрудник", "фио", "работник", "фамилия", "специалист"]},
        {"field": "accrued_amount", "label": "Начислено (руб.)", "required": True, "synonyms": ["начислено", "оклад", "начисление", "сумма начисления", "заработано"]},
        {"field": "payroll_period", "label": "Расчетный период", "required": True, "synonyms": ["период", "месяц", "расчетный месяц", "период начисления"]},
        {"field": "paid_amount", "label": "Выплачено (руб.)", "required": False, "synonyms": ["выплачено", "оплачено", "выдано", "на руки"]},
        {"field": "debt_amount", "label": "Задолженность (руб.)", "required": False, "synonyms": ["долг", "задолженность", "к выплате", "остаток"]},
        {"field": "department", "label": "Отдел / Подразделение", "required": False, "synonyms": ["отдел", "подразделение", "цех", "участок", "служба"]},
        {"field": "role_title", "label": "Должность", "required": False, "synonyms": ["должность", "профессия", "специальность"]},
    ],
    DomainCategory.OPERATIONAL_TASKS: [
        {"field": "task_title", "label": "Название поручения / задачи", "required": True, "synonyms": ["задача", "поручение", "предписание", "мероприятие", "наименование задачи", "тема"]},
        {"field": "due_date", "label": "Срок выполнения", "required": True, "synonyms": ["срок", "дедлайн", "дата выполнения", "срок исполнения", "контрольный срок"]},
        {"field": "responsible_person", "label": "Ответственный", "required": False, "synonyms": ["ответственный", "исполнитель", "куратор"]},
        {"field": "status", "label": "Статус", "required": False, "synonyms": ["статус", "состояние", "выполнение"]},
        {"field": "priority", "label": "Приоритет", "required": False, "synonyms": ["приоритет", "срочность", "важность"]},
    ],
}


class ExcelParserService:
    """
    Production-grade hybrid Excel normalization, profiling, and mapping service.
    Implements 4-step pipeline: unmerge & fill, header detection, fingerprint caching, and schema mapping.
    """

    @staticmethod
    def get_workbook_sheet_names(filepath: Path | str) -> List[str]:
        """Fast inspection of sheet names in the workbook."""
        wb = openpyxl.load_workbook(filepath, read_only=True)
        names = wb.sheetnames
        wb.close()
        return names

    @staticmethod
    def unmerge_and_normalize_sheet(ws: Worksheet) -> None:
        """
        Step 1: Unmerge all merged ranges and populate each cell with the top-left value.
        Prevents sparse matrix gaps and broken multi-row table headers.
        """
        merged_ranges = list(ws.merged_cells.ranges)
        for cell_range in merged_ranges:
            min_col, min_row, max_col, max_row = (
                cell_range.min_col,
                cell_range.min_row,
                cell_range.max_col,
                cell_range.max_row,
            )
            top_val = ws.cell(row=min_row, column=min_col).value
            ws.unmerge_cells(str(cell_range))
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    ws.cell(row=r, column=c).value = top_val

    @classmethod
    def detect_table_structure(
        cls, ws: Worksheet
    ) -> Tuple[int, List[str], List[List[Any]]]:
        """
        Detects header row, column names and separates data rows from total/summary rows.
        Scans top 15 rows for header candidates.
        """
        cls.unmerge_and_normalize_sheet(ws)

        max_row = min(ws.max_row or 1, 10000)
        max_col = ws.max_column or 1

        rows_data: List[List[Any]] = []
        for r in range(1, max_row + 1):
            row_vals = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
            rows_data.append(row_vals)

        if not rows_data:
            return 1, [], []

        # Find header row among first 15 rows based on unique column names & domain keywords
        header_row_idx = 0
        best_score = -1

        for idx, row in enumerate(rows_data[:15]):
            non_empty_strings = [
                str(v).strip()
                for v in row
                if v is not None and str(v).strip() and not isinstance(v, (int, float, datetime, date))
            ]
            unique_strings = set(non_empty_strings)
            score = len(unique_strings) * 3
            # Bonus for canonical business keywords
            keywords = ["объект", "цена", "сумма", "дата", "срок", "договор", "сотрудник", "контракт", "руб", "статья"]
            for kw in keywords:
                if any(kw in s.lower() for s in unique_strings):
                    score += 4

            if score > best_score and len(unique_strings) >= 2:
                best_score = score
                header_row_idx = idx

        # Extract column names
        raw_headers = rows_data[header_row_idx]
        headers: List[str] = []
        for col_i, val in enumerate(raw_headers):
            col_name = str(val).strip() if val is not None else ""
            if not col_name:
                col_name = f"Колонка_{col_i + 1}"
            headers.append(col_name)

        # Extract data rows following header
        data_rows: List[List[Any]] = []
        summary_keywords = ["итого", "всего", "в том числе", "общий итог", "total", "summary", "баланс"]

        for row in rows_data[header_row_idx + 1 :]:
            # Skip completely empty rows
            if not any(v is not None and str(v).strip() for v in row):
                continue

            first_val_str = str(row[0]).strip().lower() if row[0] is not None else ""
            # Exclude summary/total rows
            if any(first_val_str.startswith(kw) for kw in summary_keywords):
                continue

            data_rows.append(row)

        return header_row_idx + 1, headers, data_rows

    @staticmethod
    def calculate_headers_fingerprint(headers: List[str]) -> str:
        """
        Computes SHA-256 fingerprint from cleaned column headers.
        Used for instant 0-cost cache lookup.
        """
        normalized = [re.sub(r"\s+", " ", h.strip().lower()) for h in headers if h.strip()]
        serialized = "|".join(normalized)
        return hashlib.sha256(serialized.encode("utf-8")).hexdigest()

    @classmethod
    def infer_domain_and_mapping(
        cls, headers: List[str], sample_rows: List[List[Any]]
    ) -> Tuple[DomainCategory, float, List[Dict[str, Any]]]:
        """
        Heuristic semantic mapping engine matching columns against domain ontology.
        Returns: (suggested_domain, confidence_score, columns_mapping)
        """
        domain_scores: Dict[DomainCategory, int] = {d: 0 for d in DomainCategory}
        column_match_results: Dict[DomainCategory, List[Dict[str, Any]]] = {
            d: [] for d in DomainCategory
        }

        # Analyze each domain
        for domain, domain_fields in CANONICAL_FIELDS.items():
            used_fields = set()
            for col_idx, col_name in enumerate(headers):
                norm_col = re.sub(r"[^\w\s]", "", col_name.lower())
                best_field = None
                best_sim = 0.0

                for f_def in domain_fields:
                    f_name = f_def["field"]
                    for syn in f_def["synonyms"]:
                        # Exact or substring match
                        if syn in norm_col or norm_col in syn:
                            sim = 0.95 if syn == norm_col else 0.85
                            if sim > best_sim and f_name not in used_fields:
                                best_sim = sim
                                best_field = f_name

                sample_vals = []
                for r in sample_rows[:3]:
                    val = r[col_idx] if col_idx < len(r) and r[col_idx] is not None else ""
                    sample_vals.append(str(val))

                is_req = False
                if best_field:
                    used_fields.add(best_field)
                    domain_scores[domain] += 2
                    is_req = any(f["field"] == best_field and f["required"] for f in domain_fields)

                column_match_results[domain].append(
                    {
                        "source_column": col_name,
                        "sample_values": sample_vals,
                        "target_field": best_field,
                        "confidence": best_sim if best_field else 0.20,
                        "is_required": is_req,
                    }
                )

        # Pick best domain
        best_domain = max(domain_scores, key=lambda k: domain_scores[k])
        total_mapped = sum(1 for c in column_match_results[best_domain] if c["target_field"])
        confidence = min(0.99, max(0.50, (total_mapped / max(1, len(headers))) * 1.0))

        return best_domain, round(confidence, 2), column_match_results[best_domain]

    @staticmethod
    def clean_decimal(value: Any) -> Decimal:
        """Parses raw text/number into strict Decimal, stripping currency symbols and spaces."""
        if value is None:
            return Decimal("0.00")
        if isinstance(value, (int, float)):
            return Decimal(str(round(value, 2)))
        val_str = str(value).strip()
        if not val_str:
            return Decimal("0.00")

        # Keep only digits, commas, dots, and minus sign
        cleaned = re.sub(r"[^\d,\.-]", "", val_str)
        # Strip trailing dots/commas (e.g. from "руб." or "коп.")
        cleaned = cleaned.rstrip(".,")
        if not cleaned:
            return Decimal("0.00")

        last_comma = cleaned.rfind(",")
        last_dot = cleaned.rfind(".")
        if last_comma > last_dot:
            # Comma is the decimal separator (e.g., 12.500.000,50 or 12 500 000,50)
            int_part = cleaned[:last_comma].replace(".", "").replace(",", "")
            dec_part = cleaned[last_comma + 1 :]
            cleaned = f"{int_part}.{dec_part}"
        elif last_dot > last_comma:
            # Dot is decimal separator (e.g., 12,500,000.50 or 4.800.000.00)
            # If multiple dots, only the last is decimal separator
            int_part = cleaned[:last_dot].replace(".", "").replace(",", "")
            dec_part = cleaned[last_dot + 1 :]
            cleaned = f"{int_part}.{dec_part}"

        try:
            return Decimal(cleaned).quantize(Decimal("0.01"))
        except (InvalidOperation, ValueError):
            return Decimal("0.00")

    @staticmethod
    def clean_date(value: Any) -> Optional[date]:
        """Parses date from datetime, date, or various string formats."""
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        val_str = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y.%m.%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(val_str, fmt).date()
            except ValueError:
                continue
        return None
